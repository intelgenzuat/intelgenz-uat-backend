"""Create and load client-specific threat-actor radius scores from Excel."""

import asyncio
from collections.abc import Iterable
from pathlib import Path

import asyncpg
from openpyxl import load_workbook

from intelgenz_api.core.config import settings

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "data" / "imports" / "threat_actor_radius.xlsx"
MIGRATION_PATH = PROJECT_ROOT / "database" / "migrations" / "003_threat_actor_client_radius.sql"


def get_database_url() -> str:
    """Return the configured asyncpg connection URL."""
    database_url = settings.resolved_database_url
    if database_url is None:
        raise RuntimeError("Database settings are not configured.")
    return database_url.replace("postgresql+asyncpg://", "postgresql://", 1)


def read_radius_rows() -> list[tuple[str, str, float]]:
    """Read and validate all rows from the Radius Assessments worksheet."""
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["Radius Assessments"]
        header = next(sheet.iter_rows(max_row=1, values_only=True))
        expected_header = ("Threat Actor", "Client", "Radius")
        if header != expected_header:
            raise ValueError(f"Unexpected worksheet header: {header!r}")

        rows: list[tuple[str, str, float]] = []
        for actor_name, client_name, radius in sheet.iter_rows(min_row=2, values_only=True):
            if not isinstance(actor_name, str) or not isinstance(client_name, str):
                raise ValueError("Threat Actor and Client values must be text.")
            if not isinstance(radius, int | float) or not 0 <= radius <= 5:
                raise ValueError(f"Invalid radius for {actor_name!r}: {radius!r}")
            rows.append((actor_name.strip(), client_name.strip(), float(radius)))
    finally:
        workbook.close()

    if len(rows) != len(set((actor_name, client_name) for actor_name, client_name, _ in rows)):
        raise ValueError("The workbook contains duplicate threat actor/client pairs.")
    return rows


def find_missing_actor_names(
    rows: Iterable[tuple[str, str, float]], actor_ids_by_name: dict[str, int]
) -> list[str]:
    """Return workbook actor names that do not exist in the threat_actor table."""
    return sorted({actor_name for actor_name, _, _ in rows if actor_name not in actor_ids_by_name})


async def load_radius_assessments() -> int:
    """Apply table migration and upsert all workbook rows into the configured database."""
    radius_rows = read_radius_rows()
    connection = await asyncpg.connect(get_database_url())
    try:
        await connection.execute(MIGRATION_PATH.read_text(encoding="utf-8"))
        actor_records = await connection.fetch(
            "SELECT actor_id, canonical_name FROM public.threat_actor"
        )
        actor_ids_by_name = {
            str(record["canonical_name"]): int(record["actor_id"]) for record in actor_records
        }
        missing_actor_names = find_missing_actor_names(radius_rows, actor_ids_by_name)
        if missing_actor_names:
            preview = ", ".join(missing_actor_names[:10])
            raise ValueError(f"Workbook actors missing from threat_actor: {preview}")

        values = [
            (actor_ids_by_name[actor_name], client_name, radius)
            for actor_name, client_name, radius in radius_rows
        ]
        await connection.executemany(
            """
            INSERT INTO public.threat_actor_client_radius (actor_id, client_name, radius)
            VALUES ($1, $2, $3)
            ON CONFLICT (actor_id, client_name) DO UPDATE
            SET radius = EXCLUDED.radius, updated_at = CURRENT_TIMESTAMP
            """,
            values,
        )
        return len(values)
    finally:
        await connection.close()


if __name__ == "__main__":
    inserted_rows = asyncio.run(load_radius_assessments())
    print(f"Loaded {inserted_rows} threat actor client radius assessments.")
