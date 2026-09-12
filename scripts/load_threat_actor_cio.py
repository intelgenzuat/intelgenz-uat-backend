"""Create and load threat-actor CIO assessments from the supplied workbook."""

import asyncio
from pathlib import Path

import asyncpg
from openpyxl import load_workbook

from intelgenz_api.core.config import settings

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "data" / "imports" / "cio.xlsx"
MIGRATION_PATH = PROJECT_ROOT / "database" / "migrations" / "004_threat_actor_cio.sql"

SUMMARY_HEADER = ("Threat Actor", "Client", "Capability", "Intent", "Opportunity", "Curation")
QUESTION_HEADER = ("Assessment", "Question ID", "Question")
EVIDENCE_HEADER = (
    "Threat Actor",
    "Client",
    "Assessment",
    "Question ID",
    "Answer",
    "Actor Evidence",
    "Client Evidence",
    "Reason",
)


def get_database_url() -> str:
    """Return the configured asyncpg connection URL."""
    database_url = settings.resolved_database_url
    if database_url is None:
        raise RuntimeError("Database settings are not configured.")
    return database_url.replace("postgresql+asyncpg://", "postgresql://", 1)


def _text(value: object, field_name: str) -> str:
    """Validate and normalize a required worksheet value."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} must be a non-empty text value.")
    return value.strip()


def _client_name(value: object) -> str:
    """Use one canonical client key across CIO and radius assessment data."""
    return _text(value, "Client").upper()


def _sheet_rows(
    workbook: object, sheet_name: str, expected_header: tuple[str, ...]
) -> list[tuple[object, ...]]:
    """Read a worksheet after validating its exact data contract."""
    sheet = workbook[sheet_name]  # type: ignore[index]
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    if header != expected_header:
        raise ValueError(f"Unexpected header in {sheet_name}: {header!r}")
    return list(sheet.iter_rows(min_row=2, values_only=True))


def read_cio_workbook() -> tuple[
    list[tuple[str, str, str, str, str, str]],
    list[tuple[str, str, str]],
    list[tuple[str, str, str, str, str, str, str, str]],
]:
    """Read and validate all CIO workbook datasets."""
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        summary_rows = [
            (
                _text(actor_name, "Threat Actor"),
                _client_name(client_name),
                _text(capability, "Capability"),
                _text(intent, "Intent"),
                _text(opportunity, "Opportunity"),
                _text(curation, "Curation"),
            )
            for actor_name, client_name, capability, intent, opportunity, curation in _sheet_rows(
                workbook, "Curation Summary", SUMMARY_HEADER
            )
        ]
        question_rows = [
            (
                _text(assessment, "Assessment"),
                _text(question_id, "Question ID"),
                _text(question, "Question"),
            )
            for assessment, question_id, question in _sheet_rows(
                workbook, "Question Reference", QUESTION_HEADER
            )
        ]
        evidence_rows = [
            (
                _text(actor_name, "Threat Actor"),
                _client_name(client_name),
                _text(assessment, "Assessment"),
                _text(question_id, "Question ID"),
                _text(answer, "Answer"),
                _text(actor_evidence, "Actor Evidence"),
                _text(client_evidence, "Client Evidence"),
                _text(reason, "Reason"),
            )
            for (
                actor_name,
                client_name,
                assessment,
                question_id,
                answer,
                actor_evidence,
                client_evidence,
                reason,
            ) in _sheet_rows(workbook, "Evidence Detail", EVIDENCE_HEADER)
        ]
    finally:
        workbook.close()

    if len(summary_rows) != len({(row[0], row[1]) for row in summary_rows}):
        raise ValueError("Curation Summary contains duplicate threat actor/client pairs.")
    if len(question_rows) != len({(row[0], row[1]) for row in question_rows}):
        raise ValueError("Question Reference contains duplicate assessment/question pairs.")
    if len(evidence_rows) != len({row[:4] for row in evidence_rows}):
        raise ValueError(
            "Evidence Detail contains duplicate actor/client/assessment/question rows."
        )
    return summary_rows, question_rows, evidence_rows


async def load_cio_assessments() -> tuple[int, int, int]:
    """Apply the CIO migration and upsert all workbook rows into the configured database."""
    summary_rows, question_rows, evidence_rows = read_cio_workbook()
    connection = await asyncpg.connect(get_database_url())
    try:
        await connection.execute(MIGRATION_PATH.read_text(encoding="utf-8"))
        actor_records = await connection.fetch(
            "SELECT actor_id, canonical_name FROM public.threat_actor"
        )
        actor_ids_by_name = {
            str(record["canonical_name"]): int(record["actor_id"]) for record in actor_records
        }
        workbook_actor_names = {row[0] for row in summary_rows} | {row[0] for row in evidence_rows}
        missing_actor_names = sorted(workbook_actor_names - actor_ids_by_name.keys())
        if missing_actor_names:
            preview = ", ".join(missing_actor_names[:10])
            raise ValueError(f"Workbook actors missing from threat_actor: {preview}")

        summary_values = [
            (actor_ids_by_name[row[0]], *row[1:])
            for row in summary_rows
        ]
        evidence_values = [
            (actor_ids_by_name[row[0]], *row[1:])
            for row in evidence_rows
        ]
        async with connection.transaction():
            await connection.executemany(
                """
                INSERT INTO public.threat_actor_cio_question_reference
                    (assessment, question_id, question)
                VALUES ($1, $2, $3)
                ON CONFLICT (assessment, question_id) DO UPDATE
                SET question = EXCLUDED.question
                """,
                question_rows,
            )
            await connection.executemany(
                """
                INSERT INTO public.threat_actor_cio_curation_summary
                    (actor_id, client_name, capability, intent, opportunity, curation)
                VALUES ($1, $2, $3, $4, $5, $6)
                ON CONFLICT (actor_id, client_name) DO UPDATE
                SET capability = EXCLUDED.capability,
                    intent = EXCLUDED.intent,
                    opportunity = EXCLUDED.opportunity,
                    curation = EXCLUDED.curation
                """,
                summary_values,
            )
            await connection.executemany(
                """
                INSERT INTO public.threat_actor_cio_evidence_detail
                    (actor_id, client_name, assessment, question_id, answer,
                     actor_evidence, client_evidence, reason)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8)
                ON CONFLICT (actor_id, client_name, assessment, question_id) DO UPDATE
                SET answer = EXCLUDED.answer,
                    actor_evidence = EXCLUDED.actor_evidence,
                    client_evidence = EXCLUDED.client_evidence,
                    reason = EXCLUDED.reason
                """,
                evidence_values,
            )
        return len(summary_values), len(question_rows), len(evidence_values)
    finally:
        await connection.close()


if __name__ == "__main__":
    summary_count, question_count, evidence_count = asyncio.run(load_cio_assessments())
    print(
        "Loaded "
        f"{summary_count} CIO summaries, {question_count} question references, "
        f"and {evidence_count} evidence rows."
    )
